import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CBD5E1')
        borders.append(border)
    tblPr.append(borders)

doc = Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Normal Style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(10.5)
normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

# --- Title Header ---
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(2)
run_inst = p_title.add_run("INTERNATIONAL INFORMATION TECHNOLOGY UNIVERSITY JSC\n")
run_inst.font.bold = True
run_inst.font.size = Pt(12)
run_dept = p_title.add_run("Faculty of Business, Media and Management · Department of Information Systems\n")
run_dept.font.size = Pt(10)
run_dept.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

p_lab = doc.add_paragraph()
p_lab.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_lab.paragraph_format.space_before = Pt(6)
p_lab.paragraph_format.space_after = Pt(14)
run_lab = p_lab.add_run("REPORT ON LABORATORY WORK No. 8\n")
run_lab.font.bold = True
run_lab.font.size = Pt(15)
run_lab.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

run_topic = p_lab.add_run("The Data Access Layer: PostgreSQL Connection, JPA Entities, Associations (@OneToMany, @ManyToMany),\nSpring Data Repositories, Derived & JPQL Queries, Pagination and Sorting")
run_topic.font.size = Pt(11)
run_topic.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

# --- Metadata Table ---
table_meta = doc.add_table(rows=6, cols=2)
table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_meta)

meta_data = [
    ("Course", "RWPSF 3305 — Web Application Development with Spring Framework"),
    ("Educational Programme", "6B06105 — Information Systems, 3rd year, Semester 6"),
    ("Student", "Akishev Adilkhan (Акишев Адилхан) | Group: IS-2406 (2406IS)"),
    ("Individual Variant", "Variant 2: Attributed @ManyToMany (BookTag entity with addedAt column)"),
    ("Lecturer", "Yemberdiyeva Aknur, Senior Lector"),
    ("Git Repository & Branch", "https://github.com/adilkhanakishev/spring-lab-01/tree/lab08")
]

for i, (k, v) in enumerate(meta_data):
    row = table_meta.rows[i]
    cell_k, cell_v = row.cells[0], row.cells[1]
    cell_k.width = Inches(2.2)
    cell_v.width = Inches(4.8)
    set_cell_background(cell_k, "F8FAFC")
    set_cell_background(cell_v, "FFFFFF")
    
    p_k = cell_k.paragraphs[0]
    p_k.paragraph_format.space_before = Pt(3)
    p_k.paragraph_format.space_after = Pt(3)
    run_k = p_k.add_run(k)
    run_k.font.bold = True
    run_k.font.size = Pt(9.5)
    
    p_v = cell_v.paragraphs[0]
    p_v.paragraph_format.space_before = Pt(3)
    p_v.paragraph_format.space_after = Pt(3)
    run_v = p_v.add_run(v)
    run_v.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_heading(text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    if level == 1:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    elif level == 2:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return h

def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

def add_figure(img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption)
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

figures_dir = r"C:\Users\Asus\Desktop\spring-lab-01\docs\figures"

# Section 1
add_heading("1. Aim of the Work")
doc.add_paragraph(
    "To replace the in-memory map/list storage that carried the catalogue resource in previous laboratory works with a real database (PostgreSQL 17). "
    "The project maps persistent entities (Book, Author, Tag, and BookTag) with @OneToMany and @ManyToMany associations, answers five derived queries "
    "and two JPQL queries (including a bulk update), demonstrates and resolves the N+1 select problem using JOIN FETCH, and makes the collection "
    "REST endpoint paged and sorted with Pageable and Slice metadata."
)

add_figure(os.path.join(figures_dir, "fig1_app_started.png"), "Figure 1. Application starting against PostgreSQL with Hibernate DDL and database seeding")

# Section 2
add_heading("2. Entity Diagram and Association Cardinalities")
doc.add_paragraph(
    "• Author ↔ Book: 1:N One-to-Many association. Owning side is Book (@ManyToOne, author_id foreign key column in table book). "
    "Inverse side is Author (@OneToMany(mappedBy = 'author', cascade = CascadeType.ALL, orphanRemoval = true)).\n"
    "• Book ↔ Tag (Variant 2): N:M Many-to-Many with attributes. Upgraded to a dedicated entity BookTag mapped to table book_tag "
    "with composite primary key (@EmbeddedId BookTagId: book_id, tag_id) and custom attribute added_at (timestamp)."
)

# Section 3
add_heading("3. Database Schema Verification (PostgreSQL 17 psql)")
doc.add_paragraph("Output of psql \\dt (List of relations):")
add_figure(os.path.join(figures_dir, "fig2_tables_dt.png"), "Figure 2. PostgreSQL relations (\\dt) showing 4 created tables (author, book, book_tag, tag)")

doc.add_paragraph("Output of psql \\d book:")
add_code("""                                        Table "public.book"
   Column   |            Type             | Collation | Nullable |             Default              
------------+-----------------------------+-----------+----------+----------------------------------
 archived   | boolean                     |           | not null | 
 year       | integer                     |           | not null | 
 author_id  | bigint                      |           |          | 
 created_at | timestamp(6) with time zone |           |          | 
 id         | bigint                      |           | not null | generated by default as identity
 version    | bigint                      |           | not null | 
 genre      | character varying(20)       |           | not null | 
 title      | character varying(120)      |           | not null | 
Indexes: "book_pkey" PRIMARY KEY (id)
Foreign-key constraints: "fkklnrv3weler2ftkweewlky958" FOREIGN KEY (author_id) REFERENCES author(id)""")

doc.add_paragraph("Output of psql \\d book_tag (Variant 2 attributed join table):")
add_code("""                         Table "public.book_tag"
  Column  |            Type             | Collation | Nullable | Default 
----------+-----------------------------+-----------+----------+---------
 added_at | timestamp(6) with time zone |           | not null | 
 book_id  | bigint                      |           | not null | 
 tag_id   | bigint                      |           | not null | 
Indexes: "book_tag_pkey" PRIMARY KEY (book_id, tag_id)
Foreign-key constraints:
    "fkdrc33u5ufw8rdvajeveowgx7g" FOREIGN KEY (book_id) REFERENCES book(id)
    "fkrxw4xl05l6ns1763bq284e7m2" FOREIGN KEY (tag_id) REFERENCES tag(id)""")

# Section 4: Task 2.3 Table
add_heading("4. Task 2.3 Verification Table")
t23 = doc.add_table(rows=5, cols=3)
t23.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t23)

t23_data = [
    ("Question", "Your answer", "How you checked it"),
    ("Which tables were created?", "4 tables: author, book, book_tag, tag", "\\dt in psql and Hibernate schema DDL logs"),
    ("Where is the author foreign key?", "In table 'book' as column 'author_id' referencing author(id)", "\\d book in psql (constraint fkklnrv3weler2ftkweewlky958)"),
    ("What columns does book_tag have?", "book_id (bigint), tag_id (bigint), added_at (timestamp)", "\\d book_tag in psql (Variant 2 attributes)"),
    ("Does author have a books column?", "No. Table author only contains id, country, name", "\\d author in psql (mappedBy = 'author' has no column)")
]

for i, row_data in enumerate(t23_data):
    row = t23.rows[i]
    for j, val in enumerate(row_data):
        cell = row.cells[j]
        set_cell_background(cell, "F1F5F9" if i == 0 else "FFFFFF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(val)
        run.font.bold = (i == 0)
        run.font.size = Pt(9)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Section 5: Seed Data
add_heading("5. Database Seed Data Verification")
doc.add_paragraph("Verification of seed data in author and book tables:")
add_figure(os.path.join(figures_dir, "fig3_select_rows.png"), "Figure 3. Seed data rows in author and book tables with author_id foreign keys")

# Section 6: Task 3.3 Table
add_heading("6. Task 3.3 Query Execution Table")
t33 = doc.add_table(rows=6, cols=3)
t33.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t33)

t33_data = [
    ("Method", "Rows returned", "Number of SQL statements"),
    ("findByGenre(Genre.SCIENCE)", "3 rows (Effective Java, Java Puzzlers, Effective Java 2nd Edition)", "1 SQL query"),
    ("findByTitleContainingIgnoreCase('clean')", "3 rows (Clean Code, Clean Architecture, The Clean Coder)", "1 SQL query"),
    ("findTop3ByGenreOrderByYearDesc(Genre.SCIENCE)", "3 rows (ordered by year descending)", "1 SQL query"),
    ("search(Genre.SCIENCE, 2008) [JPQL]", "2 rows (Effective Java, Effective Java 2nd Edition)", "1 SQL query"),
    ("archiveOlderThan(2010) [bulk update]", "3 rows updated in database", "1 SQL UPDATE statement")
]

for i, row_data in enumerate(t33_data):
    row = t33.rows[i]
    for j, val in enumerate(row_data):
        cell = row.cells[j]
        set_cell_background(cell, "F1F5F9" if i == 0 else "FFFFFF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(val)
        run.font.bold = (i == 0)
        run.font.size = Pt(9)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Section 7: N+1
add_heading("7. The N+1 Exercise and Analysis (Task 3.4)")
doc.add_paragraph(
    "• Without fetch join: 4 SQL statements (1 statement to select all 8 books + 3 distinct SELECT queries for uninitialized Author proxies).\n"
    "• With fetch join (join fetch b.author): Exactly 1 SQL statement with an INNER JOIN loading both Book and Author simultaneously."
)
add_figure(os.path.join(figures_dir, "fig6_n_plus_one_sql.png"), "Figure 6. Hibernate SQL logs in IntelliJ IDEA console demonstrating query execution")

# Section 8: Pagination
add_heading("8. Pagination and Sorting (Task 4)")
doc.add_paragraph("Request: GET /api/books?page=1&size=3&sort=year,desc")
add_figure(os.path.join(figures_dir, "fig4_paged_sorted.png"), "Figure 4. Paged and sorted collection endpoint response (page=1&size=3&sort=year,desc)")

doc.add_paragraph(
    "• Page<Book> causes 2 SQL statements: one paginated SELECT with OFFSET/LIMIT and one COUNT(*) query.\n"
    "• Slice<Book> causes only 1 SQL statement (SELECT ... LIMIT (size + 1) OFFSET ...), avoiding the expensive COUNT query.\n"
    "• In invalid sort (sort=nosuchfield), Spring Data throws PropertyReferenceException, handled as 400 Bad Request ProblemDetail:"
)
add_figure(os.path.join(figures_dir, "fig5_invalid_sort_400.png"), "Figure 5. RFC 9457 ProblemDetail 400 Bad Request response for unknown sort field")

# Section 9: Variant 2
add_heading("9. Individual Variant 2: BookTag Entity with addedAt")
doc.add_paragraph(
    "Demonstration output of GET /api/demo/variant2/book-tags showing attributes stored on the join table:"
)
add_figure(os.path.join(figures_dir, "fig7_variant2_book_tags.png"), "Figure 7. Individual Variant 2 attributed Many-to-Many associations with addedAt timestamp")

# Section 10: Tests and Swagger
add_heading("10. Testing and Interactive API Documentation")
doc.add_paragraph("Execution of all 33 unit and integration tests with Maven:")
add_figure(os.path.join(figures_dir, "fig8_mvn_test_success.png"), "Figure 8. All 33 unit and integration tests passing successfully (BUILD SUCCESS)")

doc.add_paragraph("Swagger UI OpenAPI 3.1 interactive specification:")
add_figure(os.path.join(figures_dir, "fig9_swagger_ui.png"), "Figure 9. Interactive OpenAPI 3.1 Swagger UI documentation")

# Section 11: Defense Questions
add_heading("11. Key Demonstration Defence Questions")
questions = [
    ("Why can Book no longer be a record?", "JPA requires a non-final class for proxy generation, a no-arg constructor for reflection, and mutable instance fields for dirty checking."),
    ("What does @GeneratedValue(strategy = IDENTITY) do?", "Delegates PK generation to PostgreSQL's native IDENTITY/BIGSERIAL column. Value is assigned upon SQL INSERT."),
    ("Why is enum mapped with EnumType.STRING?", "Stores string names. ORDINAL stores integers (0, 1) which break if enum values are added, reordered, or removed."),
    ("What is @Version for?", "Enables Optimistic Locking. Prevents lost concurrent updates without database table locks by checking version counter during update."),
    ("Which side owns Book-Author association?", "Book owns it because Book declares @ManyToOne with @JoinColumn(name = 'author_id'). Author is inverse with mappedBy."),
    ("Why does addBook() set both sides?", "JPA only inspects the owning side. Calling only author.getBooks().add(book) would persist author_id = NULL in SQL."),
    ("Difference between join and join fetch?", "join only links tables in SQL for filtering; join fetch immediately populates associated entities, curing N+1."),
    ("Why clearAutomatically = true on @Modifying?", "Clears the Hibernate first-level cache (EntityManager) after bulk DML update so entity instances don't stay stale.")
]

for q, a in questions:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    rq = p.add_run(f"Q: {q}\n")
    rq.font.bold = True
    rq.font.size = Pt(9.5)
    ra = p.add_run(f"A: {a}")
    ra.font.size = Pt(9.5)

# Section 12: Conclusions
add_heading("12. Conclusions")
doc.add_paragraph(
    "Spring Data JPA and Hibernate ORM provide immense developer productivity by automating schema generation, synthesizing queries from "
    "method signatures, handling pagination/sorting, and maintaining consistency via dirty checking. However, developers must understand JPA lifecycle "
    "mechanics, owning association rules, and fetch strategies to prevent severe issues like N+1 queries, LazyInitializationException under "
    "open-in-view: false, and lost updates without optimistic locking."
)

out_path = r"C:\Users\Asus\Desktop\spring-lab-01\Lab08_Akishev_IS-2406_final.docx"
doc.save(out_path)
print(f"Generated {out_path} with all embedded figures successfully!")
