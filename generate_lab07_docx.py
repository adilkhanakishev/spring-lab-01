import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CBD5E1')
        borders.append(border)
    tblPr.append(borders)

doc = Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Normal Style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(10.5)
normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

# --- Title Header ---
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(2)
run_inst = p_title.add_run("INTERNATIONAL INFORMATION TECHNOLOGY UNIVERSITY JSC\n")
run_inst.font.bold = True
run_inst.font.size = Pt(13)
run_inst.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

run_dept = p_title.add_run("Faculty of Business, Media and Management · Department of Information Systems\n")
run_dept.font.size = Pt(10)
run_dept.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

run_lab = p_title.add_run("LABORATORY WORK REPORT No. 7\n")
run_lab.font.bold = True
run_lab.font.size = Pt(16)
run_lab.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

run_topic = p_title.add_run("Developing a REST API: CRUD operations on a resource, DTOs and mappers, correct HTTP status codes, error handling, documentation with OpenAPI/Swagger UI\n")
run_topic.font.italic = True
run_topic.font.size = Pt(11)
run_topic.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

# --- Metadata Table ---
tbl_meta = doc.add_table(rows=5, cols=2)
tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_meta)

meta_data = [
    [("Course:", " RWPSF 3305 — Web Application Development with Spring Framework"), 
     ("Student:", " Akishev Adilkhan")],
    [("Programme:", " 6B06105 — Information Systems, year 3, semester 6"), 
     ("Group:", " IS-2406 (2406IS)")],
    [("Lecturer:", " Yemberdiyeva Aknur, Senior Lector"), 
     ("Variant:", " 2 (Sorting and filtering: ?sort=year,desc&genre=NOVEL)")],
    [("Branch:", " lab07"), 
     ("Repository:", " https://github.com/adilkhanakishev/spring-lab-01/tree/lab07")],
    [("Pull Request:", " https://github.com/adilkhanakishev/spring-lab-01/pull/6"), 
     ("Date:", " October 2026")]
]

for row_idx, row in enumerate(tbl_meta.rows):
    row.height = Pt(20)
    for col_idx, cell in enumerate(row.cells):
        cell.width = Inches(3.5)
        set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        label, val = meta_data[row_idx][col_idx]
        r1 = p.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r2 = p.add_run(val)
        r2.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_heading_1(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    run.font.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

def add_heading_2(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    run = h.add_run(text)
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

def add_paragraph(text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    p.add_run(text)
    return p

def add_code_block(code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(7.0)
    set_cell_background(cell, "0F172A")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# 1. Aim of the Work
add_heading_1("1. Aim of the Work")
add_paragraph(
    "To turn the controller written in previous laboratory works into an enterprise-grade REST API that external teams and clients can consume reliably. "
    "The implementation establishes a clean contract on resource /api/books with five primary operations, separates request and response models via records "
    "and a standalone Spring @Component mapper, applies deliberate HTTP status codes (200, 201 + Location, 204, 404, 409), centralises RFC 9457 ProblemDetail "
    "error handling in a scoped @RestControllerAdvice, documents the API with OpenAPI 3 and Swagger UI using springdoc-openapi (v3.1.1), "
    "and fulfills Individual Variant 2 (sorting and filtering with validation)."
)

# 2. Contract Specification
add_heading_1("2. API Contract Specification")
add_paragraph("Table 1 summarizes the formal REST resource contract implemented on `/api/books`:")

tbl_contract = doc.add_table(rows=6, cols=5)
tbl_contract.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_contract)

headers = ["Method", "URI", "Request Body", "Status Codes", "Response Body"]
for i, h_text in enumerate(headers):
    cell = tbl_contract.rows[0].cells[i]
    set_cell_background(cell, "1E293B")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(h_text)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

contract_rows = [
    ("GET", "/api/books", "None", "200 OK\n400 Bad Request", "List<BookResponse>\nProblemDetail"),
    ("GET", "/api/books/{id}", "None", "200 OK\n404 Not Found", "BookResponse\nProblemDetail"),
    ("POST", "/api/books", "BookRequest (JSON)", "201 Created + Location\n400 Bad Request\n409 Conflict", "BookResponse\nProblemDetail\nProblemDetail"),
    ("PUT", "/api/books/{id}", "BookRequest (JSON)", "200 OK\n400 Bad Request\n404 Not Found\n409 Conflict", "BookResponse\nProblemDetail\nProblemDetail\nProblemDetail"),
    ("DELETE", "/api/books/{id}", "None", "204 No Content\n404 Not Found", "None\nProblemDetail")
]

for row_idx, data in enumerate(contract_rows, start=1):
    for col_idx, text in enumerate(data):
        cell = tbl_contract.rows[row_idx].cells[col_idx]
        set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.size = Pt(8.5)
        if col_idx == 0:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# 3. Code Listings
add_heading_1("3. Core Source Code Listings")

add_heading_2("Listing 1: BookRequest Record (kz.iitu.springlab.web.dto)")
add_code_block("""public record BookRequest(
        @Schema(description = "Title of the book", example = "Effective Java", maxLength = 120)
        @NotBlank @Size(max = 120) String title,

        @Schema(description = "Author of the book", example = "Joshua Bloch")
        @NotBlank String author,

        @Schema(description = "Publication year (between 1450 and 2100)", example = "2018")
        @Min(1450) @Max(2100) int year,

        @Schema(description = "Genre of the book", example = "SCIENCE")
        @NotNull Genre genre) { }""")

add_heading_2("Listing 2: BookResponse Record (kz.iitu.springlab.web.dto)")
add_code_block("""public record BookResponse(
        @Schema(description = "Unique identifier assigned by the server", example = "1")
        Long id,

        @Schema(description = "Title of the book", example = "Effective Java")
        String title,

        @Schema(description = "Author of the book", example = "Joshua Bloch")
        String author,

        @Schema(description = "Publication year", example = "2018")
        int year,

        @Schema(description = "Genre of the book", example = "SCIENCE")
        Genre genre,

        @Schema(description = "Timestamp when the book was created on the server")
        Instant createdAt) { }""")

add_heading_2("Listing 3: BookMapper Component (kz.iitu.springlab.web.dto)")
add_code_block("""@Component
public class BookMapper {
    public Book toDomain(BookRequest request) {
        return new Book(null, request.title(), request.author(),
                request.year(), request.genre(), Instant.now());
    }

    public BookResponse toResponse(Book book) {
        return new BookResponse(book.id(), book.title(), book.author(),
                book.year(), book.genre(), book.createdAt());
    }

    public List<BookResponse> toResponseList(List<Book> books) {
        return books.stream().map(this::toResponse).toList();
    }
}""")

add_heading_2("Listing 4: BookApiController (kz.iitu.springlab.web)")
add_code_block("""@Tag(name = "Books", description = "The catalogue of books")
@RestController
@RequestMapping("/api/books")
public class BookApiController {
    private final BookService service;
    private final BookMapper mapper;

    public BookApiController(BookService service, BookMapper mapper) {
        this.service = service;
        this.mapper = mapper;
    }

    @GetMapping
    public List<BookResponse> list(
            @RequestParam(required = false) String author,
            @RequestParam(required = false) Genre genre,
            @RequestParam(required = false) String sort) {
        return mapper.toResponseList(service.findAll(author, genre, sort));
    }

    @GetMapping("/{id}")
    public BookResponse find(@PathVariable long id) {
        return mapper.toResponse(service.findById(id)
                .orElseThrow(() -> new BookNotFoundException(id)));
    }

    @PostMapping
    public ResponseEntity<BookResponse> create(@Valid @RequestBody BookRequest request) {
        Book saved = service.create(mapper.toDomain(request));
        return ResponseEntity
                .created(URI.create("/api/books/" + saved.id()))
                .body(mapper.toResponse(saved));
    }

    @PutMapping("/{id}")
    public BookResponse update(@PathVariable long id, @Valid @RequestBody BookRequest request) {
        return mapper.toResponse(service.update(id, mapper.toDomain(request)));
    }

    @DeleteMapping("/{id}")
    public ResponseEntity<Void> delete(@PathVariable long id) {
        service.delete(id);
        return ResponseEntity.noContent().build();
    }
}""")

add_heading_2("Listing 5: ApiExceptionHandler (kz.iitu.springlab.web)")
add_code_block("""@RestControllerAdvice(annotations = RestController.class)
public class ApiExceptionHandler extends ResponseEntityExceptionHandler {

    @ExceptionHandler(BookNotFoundException.class)
    public ProblemDetail handleNotFound(BookNotFoundException ex) {
        ProblemDetail problem = ProblemDetail.forStatusAndDetail(
                HttpStatus.NOT_FOUND, ex.getMessage());
        problem.setType(URI.create("https://api.iitu.kz/errors/book-not-found"));
        problem.setTitle("Book not found");
        return problem;
    }

    @ExceptionHandler(DuplicateTitleException.class)
    public ProblemDetail handleDuplicate(DuplicateTitleException ex) {
        ProblemDetail problem = ProblemDetail.forStatusAndDetail(
                HttpStatus.CONFLICT, ex.getMessage());
        problem.setType(URI.create("https://api.iitu.kz/errors/duplicate-title"));
        problem.setTitle("Duplicate title");
        problem.setProperty("title", ex.getTitle());
        return problem;
    }

    @ExceptionHandler(InvalidSortFieldException.class)
    public ProblemDetail handleInvalidSortField(InvalidSortFieldException ex) {
        ProblemDetail problem = ProblemDetail.forStatusAndDetail(
                HttpStatus.BAD_REQUEST, ex.getMessage());
        problem.setType(URI.create("https://api.iitu.kz/errors/invalid-sort-field"));
        problem.setTitle("Invalid sort field");
        problem.setProperty("field", ex.getField());
        return problem;
    }
}""")

# 4. Experimental Results - Task 2.3
add_heading_1("4. Experimental Results: Task 2.3 Completed Table")
add_paragraph("Table 2 records the actual observed status and response content from the running server:")

tbl_23 = doc.add_table(rows=7, cols=3)
tbl_23.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_23)

h_23 = ["Request Command", "Status", "Observed Response Headers & Body"]
for i, h_text in enumerate(h_23):
    cell = tbl_23.rows[0].cells[i]
    set_cell_background(cell, "1E293B")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(h_text)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

data_23 = [
    ("curl -i localhost:8080/api/books", "200 OK", "Content-Type: application/json\nJSON array of 3 books: Effective Java, Clean Code, Spring in Action with DTO fields id, title, author, year, genre, createdAt."),
    ("curl -i localhost:8080/api/books/1", "200 OK", "Content-Type: application/json\nJSON object for book 1: {\"id\":1,\"title\":\"Effective Java\",\"author\":\"Joshua Bloch\",\"year\":2018,\"genre\":\"SCIENCE\",\"createdAt\":\"...\"}."),
    ("curl -i localhost:8080/api/books/999", "404 Not Found", "Content-Type: application/problem+json\n{\"type\":\"https://api.iitu.kz/errors/book-not-found\",\"title\":\"Book not found\",\"status\":404,\"detail\":\"Book 999 not found\",\"instance\":\"/api/books/999\"}."),
    ("curl -i -X POST -H \"Content-Type: application/json\" -d '...' .../api/books", "201 Created", "Location: /api/books/4\nContent-Type: application/json\nBody contains created book with server-assigned ID 4 and current timestamp."),
    ("curl -i -X DELETE localhost:8080/api/books/1", "204 No Content", "No body (Content-Length: 0). The resource was permanently deleted from memory storage."),
    ("Repeated DELETE localhost:8080/api/books/1", "404 Not Found", "Content-Type: application/problem+json\n{\"detail\":\"Book 1 not found\",\"instance\":\"/api/books/1\",\"status\":404,\"title\":\"Book not found\",...}")
]

for row_idx, row_vals in enumerate(data_23, start=1):
    for col_idx, text in enumerate(row_vals):
        cell = tbl_23.rows[row_idx].cells[col_idx]
        set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.size = Pt(8.5)
        if col_idx == 1:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# 5. Error Routes - Task 3.3
add_heading_1("5. Error Handling: Task 3.3 Completed Table")

tbl_33 = doc.add_table(rows=4, cols=3)
tbl_33.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_33)

h_33 = ["Request", "Status and Content-Type", "Which Handler Executed"]
for i, h_text in enumerate(h_33):
    cell = tbl_33.rows[0].cells[i]
    set_cell_background(cell, "1E293B")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(h_text)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

data_33 = [
    ("GET /api/books/999", "404 Not Found\napplication/problem+json", "ApiExceptionHandler.handleNotFound(BookNotFoundException)"),
    ("POST /api/books with empty title", "400 Bad Request\napplication/problem+json", "ResponseEntityExceptionHandler.handleMethodArgumentNotValid (Base Class)"),
    ("POST /api/books duplicate book", "409 Conflict\napplication/problem+json", "ApiExceptionHandler.handleDuplicate(DuplicateTitleException)")
]

for row_idx, row_vals in enumerate(data_33, start=1):
    for col_idx, text in enumerate(row_vals):
        cell = tbl_33.rows[row_idx].cells[col_idx]
        set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.size = Pt(8.5)
        if col_idx == 1:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

add_heading_2("Full RFC 9457 ProblemDetail Body (GET /api/books/999):")
add_code_block("""{
  "type": "https://api.iitu.kz/errors/book-not-found",
  "title": "Book not found",
  "status": 404,
  "detail": "Book 999 not found",
  "instance": "/api/books/999"
}""")

# 6. Screenshots
add_heading_1("6. Demonstration Evidence and Screenshots")

if os.path.exists("docs/figures/fig1_post_201_location.png"):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture("docs/figures/fig1_post_201_location.png", width=Inches(6.5))
    p_cap1 = doc.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap1 = p_cap1.add_run("Figure 1: Successful POST 201 Created with Location Header and Verification via GET")
    r_cap1.font.italic = True
    r_cap1.font.size = Pt(8.5)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

if os.path.exists("docs/figures/fig2_swagger_ui_404.png"):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture("docs/figures/fig2_swagger_ui_404.png", width=Inches(6.5))
    p_cap2 = doc.add_paragraph()
    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap2 = p_cap2.add_run("Figure 2: Swagger UI Interactive GET /api/books/{id} Execution and 404 ProblemDetail")
    r_cap2.font.italic = True
    r_cap2.font.size = Pt(8.5)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

if os.path.exists("docs/figures/fig3_variant2_sorting.png"):
    p_img3 = doc.add_paragraph()
    p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture("docs/figures/fig3_variant2_sorting.png", width=Inches(6.5))
    p_cap3 = doc.add_paragraph()
    p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap3 = p_cap3.add_run("Figure 3: Individual Variant 2 Execution — Sorting, Filtering, and Validation")
    r_cap3.font.italic = True
    r_cap3.font.size = Pt(8.5)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

if os.path.exists("docs/figures/fig4_github_pr6.png"):
    p_img4 = doc.add_paragraph()
    p_img4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture("docs/figures/fig4_github_pr6.png", width=Inches(6.5))
    p_cap4 = doc.add_paragraph()
    p_cap4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap4 = p_cap4.add_run("Figure 4: Official GitHub Pull Request #6 (lab07 -> main)")
    r_cap4.font.italic = True
    r_cap4.font.size = Pt(8.5)

# 7. Individual Variant 2
add_heading_1("7. Individual Assignment: Variant 2 (Sorting & Filtering)")
add_paragraph(
    "Variant 2 requires collection sorting and filtering (?sort=year,desc&genre=NOVEL), with an unknown sort field "
    "returning 400 Bad Request with an RFC 9457 ProblemDetail payload.",
    bold_prefix="Specification: "
)
add_paragraph(
    "Implemented in BookService.findAll(author, genre, sort) by creating dynamic Comparators across supported fields "
    "(year, title, author, id, genre, createdAt) with asc/desc direction handling. When an unmapped sort field (e.g. isbn) is supplied, "
    "InvalidSortFieldException is thrown and caught by ApiExceptionHandler, returning 400 Bad Request with custom extension property field.",
    bold_prefix="Architecture: "
)

# 8. Questions and Answers
add_heading_1("8. Questions Asked During Demonstration (Section 10)")

qa_list = [
    ("Why is the URI a noun? What would be wrong with /api/getBooks?",
     "REST publishes resources (nouns), and HTTP verbs define actions. /api/getBooks is an RPC antipattern that explodes the URI namespace and duplicates protocol verbs."),
    
    ("Which of your endpoints are safe and which are idempotent?",
     "Safe: GET /api/books, GET /api/books/{id}. Idempotent: GET, PUT /api/books/{id}, and DELETE /api/books/{id}. POST is neither safe nor idempotent."),
    
    ("Why does creation return 201 and not 200? What does the Location header give the client?",
     "201 Created explicitly confirms resource instantiation. The Location header (/api/books/4) provides the canonical URI of the new entity for direct retrieval without parsing the response body."),
    
    ("Why does deletion return 204 and not 200 with a body?",
     "204 No Content signals complete processing where no remaining entity exists to be returned. Returning 200 with arbitrary JSON wastes bandwidth and complicates client parsing."),
    
    ("What does your API answer to a second DELETE of the same identifier, and why that answer?",
     "Our API returns 404 Not Found (ProblemDetail) because the resource ceased to exist. Returning 204 is also valid under strict RFC idempotency; what is inadmissible is 500."),
    
    ("What is the difference between PUT and PATCH for a field missing from the body?",
     "PUT replaces the entire entity (missing fields are reset to defaults/null). PATCH modifies only the fields explicitly included in the request body."),
    
    ("What is a DTO? Give three things that would go wrong if you returned the entity.",
     "A Data Transfer Object separates API contracts from storage models. Exposing entities causes: 1) Mass assignment security vulnerabilities, 2) Breaking API changes upon DB column renames, 3) Accidental data leakage and circular JSON recursion."),
    
    ("Why are the request DTO and the response DTO separate types in your code?",
     "BookRequest enforces input validation and excludes server-assigned fields (id, createdAt). BookResponse includes server metadata and matches what the API contract promises."),
    
    ("Where does your mapper live and why is it not in the controller?",
     "BookMapper is a standalone Spring @Component in kz.iitu.springlab.web.dto. Keeping mapping out of the controller honors Single Responsibility, declutters routing, and enables independent unit testing."),
    
    ("Which status code answers an unparsable body, and which answers a body that is valid JSON but breaks a rule?",
     "Both return 400 Bad Request: malformed JSON triggers HttpMessageNotReadableException, while rule-breaking JSON triggers MethodArgumentNotValidException."),
    
    ("When do you return 409? Why not 400?",
     "409 Conflict is returned when the request syntax is valid, but server state prevents execution (e.g., duplicate title and author). 400 indicates client syntax errors."),
    
    ("Which RFC defines ProblemDetail? Name its five standard fields.",
     "RFC 9457 defines ProblemDetail. Its five standard fields are: type (URI), title (String), status (int), detail (String), and instance (URI)."),
    
    ("Which media type does an error response carry, and who sets it?",
     "It carries application/problem+json, set automatically by Spring's ProblemDetail message converters and ResponseEntityExceptionHandler."),
    
    ("What does extending ResponseEntityExceptionHandler give you that you did not write?",
     "Automatic RFC 9457 conversion for standard Spring MVC exceptions (type mismatches, missing parameters, validation errors, unsupported HTTP methods) with proper HTTP status codes."),
    
    ("How is your advice scoped? What would happen to the pages of lab 6 without the scoping?",
     "Scoped via @RestControllerAdvice(annotations = RestController.class). Without scoping, browser navigation to Lab 6 pages (/books/999) would receive raw JSON instead of the Thymeleaf HTML error view."),
    
    ("Why must a stack trace never appear in a response body?",
     "Stack traces expose internal architecture, packages, and library versions to attackers, creating security vulnerabilities. Clients cannot parse stack traces, which belong in server logs."),
    
    ("Which parts of the OpenAPI document were generated without any annotation from you?",
     "All endpoint routes, HTTP verbs, JSON media types, request/response DTO property types via reflection, and validation constraints (maxLength: 120, min: 1450, max: 2100, required)."),
    
    ("Where did the maxLength of title in the schema come from?",
     "From the Jakarta validation annotation @Size(max = 120) on the title component of BookRequest, automatically inspected by SpringDoc."),
    
    ("At which URLs are the document and the Swagger UI page served by default?",
     "OpenAPI JSON specification: /v3/api-docs. Swagger UI HTML interface: /swagger-ui.html (redirecting to /swagger-ui/index.html)."),
    
    ("Why should Swagger UI not be left open on a public server, and how would you close it?",
     "Leaving it exposed allows attackers to explore all endpoints, parameters, and execute live attacks. It is disabled via properties: springdoc.api-docs.enabled=false and springdoc.swagger-ui.enabled=false.")
]

for q, a in qa_list:
    add_paragraph(q, bold_prefix="Q: ")
    p_ans = add_paragraph(a, bold_prefix="A: ", space_after=6)
    p_ans.runs[1].font.color.rgb = RGBColor(0x33, 0x41, 0x55)

# 9. Conclusions
add_heading_1("9. Conclusions")
add_paragraph(
    "Implementing a dedicated DTO and mapper layer introduces a minimal initial cost in boilerplate code and object allocation. "
    "However, what it buys for an enterprise system is substantial: it guarantees strict architectural decoupling between external API contracts "
    "and internal persistence entities, completely eliminates mass assignment vulnerabilities, prevents accidental leakage of sensitive attributes, "
    "and provides predictable, standardized communication using RFC 9457 ProblemDetail payloads and interactive OpenAPI/Swagger documentation."
)

output_filename = "Lab07_Akishev_IS-2406.docx"
doc.save(output_filename)
print(f"Successfully generated {output_filename}")
